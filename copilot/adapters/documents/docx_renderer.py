"""ATS-friendly DOCX resume: single column, no tables or images, standard headings.
Indian convention: 1–2 pages, no photo by default."""

from __future__ import annotations

from pathlib import Path

from copilot.adapters.documents.markdown_renderer import resume_sections, safe_filename
from copilot.domain.models import CandidateProfile, TailoredResume
from copilot.ports.documents import RenderedDocument


class DocxResumeRenderer:
    def render_resume(
        self, profile: CandidateProfile, tailored: TailoredResume, output_dir: Path
    ) -> RenderedDocument:
        from docx import Document  # noqa: PLC0415 - optional dependency
        from docx.shared import Pt  # noqa: PLC0415

        output_dir.mkdir(parents=True, exist_ok=True)
        doc = Document()
        doc.styles["Normal"].font.name = "Calibri"
        doc.styles["Normal"].font.size = Pt(10.5)
        for i, (heading, body) in enumerate(resume_sections(profile, tailored)):
            doc.add_heading(heading, level=0 if i == 0 else 1)
            for line in body:
                style = "List Bullet" if i >= 2 else None
                doc.add_paragraph(line, style=style)
        path = output_dir / f"{safe_filename(profile.full_name, tailored.job_id)}.docx"
        doc.save(str(path))
        return RenderedDocument(path=path, format="docx")
